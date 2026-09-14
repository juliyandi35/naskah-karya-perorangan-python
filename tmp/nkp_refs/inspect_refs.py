from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def twips(value):
    return None if value is None else round(value.twips)


def inspect(path: Path) -> dict:
    doc = Document(path)
    data = {
        "file": str(path),
        "sections": [],
        "styles": {},
        "paragraphs": [],
        "tables": [],
    }
    for i, section in enumerate(doc.sections, 1):
        data["sections"].append({
            "index": i,
            "width_twips": twips(section.page_width),
            "height_twips": twips(section.page_height),
            "top_margin_twips": twips(section.top_margin),
            "bottom_margin_twips": twips(section.bottom_margin),
            "left_margin_twips": twips(section.left_margin),
            "right_margin_twips": twips(section.right_margin),
            "header_distance_twips": twips(section.header_distance),
            "footer_distance_twips": twips(section.footer_distance),
            "start_type": str(section.start_type),
        })
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        pf = style.paragraph_format
        font = style.font
        data["styles"][name] = {
            "font": font.name,
            "size_pt": None if font.size is None else font.size.pt,
            "bold": font.bold,
            "italic": font.italic,
            "alignment": str(pf.alignment),
            "line_spacing": str(pf.line_spacing),
            "space_before_pt": None if pf.space_before is None else pf.space_before.pt,
            "space_after_pt": None if pf.space_after is None else pf.space_after.pt,
            "first_line_indent_twips": twips(pf.first_line_indent),
            "left_indent_twips": twips(pf.left_indent),
        }
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.replace("\t", "<TAB>").strip()
        if not text:
            continue
        pf = paragraph.paragraph_format
        run_info = []
        for run in paragraph.runs:
            if run.text.strip():
                run_info.append({
                    "text": run.text,
                    "font": run.font.name,
                    "size_pt": None if run.font.size is None else run.font.size.pt,
                    "bold": run.bold,
                    "italic": run.italic,
                    "underline": run.underline,
                })
        data["paragraphs"].append({
            "index": i,
            "style": paragraph.style.name if paragraph.style else None,
            "text": text,
            "alignment": str(paragraph.alignment),
            "first_line_indent_twips": twips(pf.first_line_indent),
            "left_indent_twips": twips(pf.left_indent),
            "right_indent_twips": twips(pf.right_indent),
            "space_before_pt": None if pf.space_before is None else pf.space_before.pt,
            "space_after_pt": None if pf.space_after is None else pf.space_after.pt,
            "line_spacing": str(pf.line_spacing),
            "page_break_before": pf.page_break_before,
            "keep_with_next": pf.keep_with_next,
            "runs": run_info,
        })
    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip().replace("\n", " | ") for cell in row.cells])
        data["tables"].append({"index": ti, "style": table.style.name if table.style else None, "rows": rows})
    return data


for arg in sys.argv[1:]:
    path = Path(arg).resolve()
    out = path.with_suffix(".inspection.json")
    out.write_text(json.dumps(inspect(path), ensure_ascii=False, indent=2), encoding="utf-8")
    print(out)
